import os
from pathlib import Path
import sys
import zipfile
from PIL import Image, ImageDraw, ImageFont
from reportlab.pdfgen import canvas
from docx import Document
from openpyxl import Workbook
from pypdf import PdfReader, PdfWriter

root = Path(sys.argv[1]); root.mkdir(mode=0o700, parents=True, exist_ok=True)
image = Image.new('RGB', (900, 360), 'white'); draw = ImageDraw.Draw(image)
font = ImageFont.load_default(size=70)
draw.text((45, 35), 'VISION-731', fill='navy', font=font)
draw.ellipse((45, 160, 175, 290), fill='red')
draw.rectangle((260, 160, 600, 290), fill='yellow')
for ext, form in [('png','PNG'),('jpg','JPEG'),('webp','WEBP'),('gif','GIF')]: image.save(root / ('sample.' + ext), format=form)
c = canvas.Canvas(str(root/'sample.pdf')); c.drawString(72, 720, 'PDF-482'); c.drawString(72, 690, 'Synthetic bounded document input'); c.save()
scan = Image.new('RGB',(900,360),'white'); ImageDraw.Draw(scan).text((45,80),'SCAN-926',fill='green',font=font)
scan.save(root/'scan.png')
c = canvas.Canvas(str(root/'scan.pdf')); c.drawImage(str(root/'scan.png'),45,400,width=500,height=200); c.save()
document = Document(); document.add_paragraph('DOCX-593'); document.add_table(rows=1, cols=2).rows[0].cells[0].text = 'Table content'
document.save(root/'sample.docx')
with zipfile.ZipFile(root/'sample.docx') as source, zipfile.ZipFile(root/'unsafe.docx','w') as output:
    for member in source.infolist(): output.writestr(member, source.read(member.filename))
    output.writestr('../outside.txt','must never be extracted')
book = Workbook(); book.active['A1'] = 'XLSX-264'; book.active['A2'] = 42; book.active['B2'] = '=1+1'; book.save(root/'sample.xlsx')
(root/'sample.txt').write_text('TEXT-815\n合成文本，不含私人信息。', encoding='utf-8')
(root/'large.txt').write_text('x'*60000, encoding='utf-8')
writer = PdfWriter(); writer.add_page(PdfReader(root/'sample.pdf').pages[0]); writer.encrypt('synthetic-only')
with (root/'encrypted.pdf').open('wb') as f: writer.write(f)
for path in root.iterdir(): os.chmod(path, 0o600)
